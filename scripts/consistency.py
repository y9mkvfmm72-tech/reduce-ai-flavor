# -*- coding: utf-8 -*-
"""Does the prose agree with the document's own tables and figures?

    python consistency.py 定稿.docx
    python consistency.py 定稿.docx --max-claims 15

Read-only worksheet, not a checker. Six lists:

  [1] prose claims (compare / constant / total language)  -- verify against [2]
      + mechanism sentences (a cause offered for the numbers)
  [2] near-constant numeric rows, with their relative spread
  [3] numbers in the prose that no table contains          -- derived? borrowed? unsupported?
  [4] naming sentences ("here it is written X") and one-off Greek symbols
  [5] table inventory + every "of Table N" cross-reference
  [6] artefacts: figures / captions the text needs but the file lacks,
      placeholder algebra (x%, 4x%), and tables cited but not captioned

Why this script exists: on the case that produced this skill, the style audit
passed but three sentences still contradicted the document's own tables
("nearly the same" against a 6.2% spread, "integrating gives the total" against
a 70%-coverage window, "K(lambda) of Table 1" where Table 1 has no K). A reader
with a calculator finds those faster than any style tell, and no delete-only
pass can fix them: the sentence itself is wrong. The next round added two more
of the same kind - a mechanism ("largest where the transfer function changes
quickly") that the tables contradict, and a document that described curves it
never showed.
"""
import argparse
import io
import os
import re
import sys

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    # Offline / locked-down Python: stdlib reader, same surface.
    from docxfallback import Document, qn

INVISIBLE = "\u2060\u200b\u200c\u200d\u200e\u200f\ufeff"
GREEK_RE = re.compile(r"[\u0391-\u03a9\u03b1-\u03c9]")
NUM = re.compile(r"[+\-\u2212]?\d+(?:[.,]\d+)?(?:\s*[\u00d7x]\s*10\s*\^?\s*[\u2212\-]?\d+)?")

STRONG = [r"\bnearly\b", r"\balmost\b", r"\bidentical\b", r"\bthe same\b",
          r"\bconstant\b", r"\bdoes not (?:change|vary)\b", r"\bno (?:larger|greater) than\b"]
WEAK = [r"\bwithin\b", r"\bagree", r"\bconsistent\b", r"\bmean\b", r"\baverage\b",
        r"\btotal\b", r"\bclose to\b"]
NAMING = [r"\bwritten as\b", r"\bhere it is written\b", r"\bdenoted\b", r"\blabell?ed\b",
          r"\blabels?\b", r"\breferred to as\b", r"\bcalled\b",
          "记为", "记作", "写作", "称为"]
CAPTION = re.compile(r"^\s*(?:Table\s+|表\s*)([0-9]+|[A-Z]\d+)[.．:：]?", re.I)


def caption_number(text):
    m = CAPTION.match(text or "")
    return m.group(1) if m else None

# A cause asserted without a measured basis. Check each against the tables.
MECHANISM = [
    r"\bis the signature of\b",
    r"\b(?:this|the offset|the spread|the scatter) points? to\b",
    r"\bis dominated by\b",
    r"\bcan be (?:attributed|explained|accounted for) by\b",
    r"\bis (?:therefore )?due to\b",
    r"\bbecause\b[^.]{0,90}\d",
    r"\b(?:the|this) (?:likely|probable) (?:cause|reason|explanation)\b",
    r"\barises from\b",
]
# Symbols the text never instantiates (x%, 4x%, "for some x").
PLACEHOLDER = [
    r"\bx\s*%", r"\b\d+x\s*%", r"\bfor (?:every|each) x\b", r"\bsome x\b",
    r"\bx\s*times\b", r"\bproportional to x\b", r"\barbitrary (?:x|n)\b",
]
FIG_REF = re.compile(r"(?:\bFig(?:ure)?\.?\s*|图\s*)(\d+)", re.I)
CAP_FIG = re.compile(r"^\s*(?:Fig(?:ure)?\.?\s*|图\s*)(\d+)[.．:：]?", re.I)
TAB_REF = re.compile(r"(?:\bTables?\s+|表\s*)([0-9]+|[A-Z]\d+)", re.I)

RENAME = re.compile(
    r"(?:here it is written|written as|denoted|labelled|labeled|referred to as)"
    r"\s*([\u0391-\u03a9\u03b1-\u03c9])", re.I)


def clean(t):
    for c in INVISIBLE:
        t = t.replace(c, "")
    return t


def num_of(tok):
    tok = tok.replace("\u2212", "-").replace("+", "").strip()
    tok = re.split(r"[\u00d7x]", tok)[0].strip().replace(",", "")
    try:
        return float(tok)
    except ValueError:
        return None


def sentences(text):
    parts = re.split(r"(?<=[.!?])\s+|(?<=[。！？])", text)
    return [s.strip() for s in parts if s.strip()]


def load(path):
    if os.path.splitext(path)[1].lower() != ".docx":
        sys.exit("this script needs a .docx: the tables are the evidence")
    if Document is None:
        sys.exit("python-docx is required")
    d = Document(path)
    paras = [clean(" ".join(p.text.split())) for p in d.paragraphs if p.text.strip()]

    body = d.element.body
    tbl_elems = [el for el in body.iterchildren() if el.tag == qn("w:tbl")]
    tables = []
    for i, el in enumerate(tbl_elems):
        t = d.tables[i]
        # 1-row borderless tables are equation-layout tables, not data tables.
        if len(t.rows) == 1 and sum(1 for c in t.rows[0].cells if c.text.strip()) <= 1:
            continue
        cap = None
        # Captions may sit above (Chinese reports) or below (lab-notebook style).
        sib, steps = el.getprevious(), 0
        while sib is not None and steps < 4:
            if sib.tag != qn("w:p"):
                break
            cap = caption_number("".join(sib.itertext()).strip())
            if cap:
                break
            steps += 1
            sib = sib.getprevious()
        if cap is None:
            sib, steps = el.getnext(), 0
            while sib is not None and steps < 4:
                if sib.tag != qn("w:p"):
                    break
                cap = caption_number("".join(sib.itertext()).strip())
                if cap:
                    break
                steps += 1
                sib = sib.getnext()
        grid = [[clean(c.text).strip() for c in r.cells] for r in t.rows]
        tables.append((cap, grid))
    return d, paras, tables


def numeric_rows(tables):
    rows = []
    for cap, grid in tables:
        drop_mean = bool(grid) and grid[0] and \
            grid[0][-1].strip().lower() in ("mean", "average", "平均")
        for cells in grid:
            nums = [v for v in (num_of(c) for c in cells[1:]) if v is not None]
            if drop_mean and len(nums) == len(cells) - 1 and len(nums) > 1:
                nums = nums[:-1]
            if len(nums) >= 3:
                rows.append((cap, cells[0], nums))
    return rows


def main(path, max_claims):
    d, paras, tables = load(path)
    if not tables:
        print("no tables in %s - this script needs tabular ground truth; "
              "use flavor_audit.py instead" % os.path.basename(path))
        return 0
    body = " ".join(paras)
    sents = sentences(body)
    rows = numeric_rows(tables)
    tab_all = []
    for _, grid in tables:
        for cells in grid:
            for c in cells:
                v = num_of(c)
                if v is not None:
                    tab_all.append(v)
    for p in paras:                       # constants often live in captions
        if CAPTION.match(p):
            for tok in NUM.findall(p):
                v = num_of(tok)
                if v is not None:
                    tab_all.append(v)

    print("file   : %s" % os.path.basename(path))
    print("tables : %d   numeric rows: %d   prose sentences: %d"
          % (len(tables), len(rows), len(sents)))

    # ---- 1 claims ---------------------------------------------------------
    claims = []
    for s in sents:
        if len(s) < 20:
            continue
        strong = any(re.search(k, s, re.I) for k in STRONG)
        weak = any(re.search(k, s, re.I) for k in WEAK) and \
            re.search(r"\d+\.\d+|\d+\s*%|\d+\s*[\u00d7x]\s*10", s)
        if strong or weak:
            claims.append(s)
    print("\n[1] prose claims (compare / generalise) -- verify each against [2]")
    for s in claims[:max_claims]:
        print("    - %s" % s[:170])
    if len(claims) > max_claims:
        print("    ... %d more" % (len(claims) - max_claims))
    print("    %s" % ("PASS (none)" if not claims else
                     "REVIEW: %d sentences make a quantitative claim" % len(claims)))

    mech = [s for s in sents if any(re.search(rx, s, re.I) for rx in MECHANISM)]
    if mech:
        print("\n    mechanism sentences -- the tables must support the cause:")
        for s in mech[:max_claims]:
            print("    * %s" % s[:170])

    # ---- 2 near-constant rows --------------------------------------------
    flat, other = [], 0
    for cap, label, nums in rows:
        mean = sum(nums) / len(nums)
        spread = (max(nums) - min(nums)) / abs(mean) * 100 if mean else float("inf")
        if spread <= 10:
            flat.append((cap, label, nums, mean, spread))
        else:
            other += 1
    print("\n[2] near-constant numeric rows (spread = (max-min)/|mean|)")
    for cap, label, nums, mean, spread in flat:
        tag = ("Table %d" % cap) if cap else "cover"
        vals = " ".join("%g" % v for v in nums[:6])
        print("    %-8s %-24s n=%-2d mean=%-9.4g spread=%5.1f%%  %s"
              % (tag, label[:24], len(nums), mean, spread, vals))
    print("    %d other numeric rows vary by more than 10%% (spectra, grids) - not listed"
          % other)
    print("    a claim of 'constant / the same / within x%' must survive the spread above")

    # ---- 3 numbers not in tables -----------------------------------------
    missing, seen = [], []
    for tok in NUM.findall(body):
        v = num_of(tok)
        if v is None or re.fullmatch(r"[+\-\u2212]?\d{1,2}", tok.strip()):
            continue
        if any(abs(v - t) <= max(abs(t) * 0.002, 1e-12) for t in tab_all):
            continue
        if any(abs(v - x) <= max(abs(x) * 0.002, 1e-12) for x in seen):
            continue
        seen.append(v)
        missing.append(tok.strip())
    print("\n[3] numbers in the prose that no table contains")
    print("    %s" % (", ".join(missing[:30]) if missing else "(none)"))
    if missing:
        print("    derived from the tables? borrowed from the manual? unsupported? "
              "an unexplained one is a claim, not a measurement")

    # ---- 4 symbols --------------------------------------------------------
    names = [s for s in sents if any(re.search(k, s, re.I) for k in NAMING)]
    print("\n[4] naming sentences and one-off Greek symbols")
    for s in names:
        print("    name: %s" % s[:150])
        m = RENAME.search(s)
        if m:
            new = m.group(1)
            before = s[:m.start()]
            olds = GREEK_RE.findall(before)
            if olds and olds[-1] != new:
                old = olds[-1]
                print("      RENAME %s -> %s; text still uses %s x%d and %s x%d"
                      % (old, new, old, body.count(old), new, body.count(new)))
                if body.count(old) > body.count(new):
                    print("      the sentence is a leftover: %s still dominates"
                          % old)
    counts = {}
    for ch in GREEK_RE.findall(body):
        counts[ch] = counts.get(ch, 0) + 1
    rare = ["%s x%d" % (c, n) for c, n in sorted(counts.items(), key=lambda kv: kv[1])
            if n <= 2]
    print("    Greek in prose: %s" % ("  ".join(
        "%s x%d" % (c, n) for c, n in sorted(counts.items(), key=lambda kv: -kv[1]))
        or "(none)"))
    print("    one-or-two uses: %s" % (", ".join(rare) if rare else "(none)"))

    # ---- 5 inventory + cross-references ----------------------------------
    print("\n[5] table inventory -- check every 'of Table N' against this")
    by_num = {}
    for cap, grid in tables:
        head = " | ".join(c for c in (grid[0] if grid else []) if c)
        by_num[cap] = " ".join(" ".join(r) for r in grid)
        print("    %-8s %s" % (("Table %s" % cap) if cap else "cover", head[:110]))
    flags = []
    for s in sents:
        for m in re.finditer(r"(?:of\s+Table|表)\s*([0-9]+|[A-Z]\d+)", s):
            n = m.group(1)
            frag = s[max(0, m.start() - 45):m.start()]
            greek = set(GREEK_RE.findall(frag))
            miss = [g for g in greek if g not in by_num.get(n, "")]
            if miss:
                flags.append("Table %s is cited for %s, but %s does not appear in it: %s"
                             % (n, "".join(miss), "".join(miss), s[:120]))
    for f in flags:
        print("    CROSS-REF: %s" % f)
    if not flags:
        print("    no symbol-level cross-reference mismatch found")

    # ---- 6 artefacts and placeholders ------------------------------------
    print("\n[6] artefacts and placeholders")
    images = len(d.inline_shapes)
    try:
        drawings = len(d.element.body.findall(".//" + qn("w:drawing")))
    except Exception:
        drawings = images
    n_images = max(images, drawings)
    fig_refs = sorted({int(m.group(1)) for m in FIG_REF.finditer(body)})
    fig_caps = sorted({int(m.group(1)) for p in paras for m in CAP_FIG.finditer(p)})
    tab_refs = sorted({m.group(1) for m in TAB_REF.finditer(body)})
    caps = sorted({c for c, _ in tables if c})
    print("    images %d | figure refs %s | figure captions %s"
          % (n_images, fig_refs or "-", fig_caps or "-"))
    print("    tables %d (%d captioned) | table refs %s | captions %s"
          % (len(tables), len(caps), tab_refs or "-", caps or "-"))
    if fig_refs and n_images == 0:
        print("    MISSING-FIGURE: the prose cites Figure %s, the file has no image"
              % ", ".join(map(str, fig_refs)))
    if fig_caps and n_images == 0:
        print("    MISSING-FIGURE: a figure caption exists, the file has no image")
    curve_hits = len(re.findall(r"\b(?:curve|curves|plot|plotted|graph|spectrum|spectra)\b",
                                body, re.I))
    if n_images == 0 and not fig_caps and curve_hits >= 3:
        print("    MISSING-FIGURE: no image in the file, yet the prose describes "
              "curves/plots %d times" % curve_hits)
    ghost_tab = [c for c in tab_refs if c not in caps]
    if ghost_tab and len(tables) > 0:
        print("    MISSING-TABLE: prose cites Table %s, captions only %s"
              % (", ".join(map(str, ghost_tab)), caps or "none"))
    ph = [s for s in sents if any(re.search(rx, s, re.I) for rx in PLACEHOLDER)]
    for s in ph:
        print("    PLACEHOLDER: %s" % s[:150])
    if not (fig_refs and n_images == 0) and not ghost_tab and not ph:
        print("    PASS: no missing figure/table, no placeholder algebra")

    print("\nreading: [1] vs [2] is where contradictions show up; [3]/[4] are "
          "residues; [5]/[6] are the reference checks. A contradiction is not a "
          "style problem: deleting adjectives will not fix it, the sentence must "
          "be rewritten or cut. A missing figure is cheapest to fix by drawing "
          "it from the same data pipeline as the tables.")
    return 0


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                  errors="replace")
    ap = argparse.ArgumentParser()
    ap.add_argument("mine", help="document with tables (.docx)")
    ap.add_argument("--max-claims", type=int, default=25)
    a = ap.parse_args()
    sys.exit(main(a.mine, a.max_claims))
