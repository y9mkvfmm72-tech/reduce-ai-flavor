# -*- coding: utf-8 -*-
"""How much of my document is somebody else's text?

    python overlap.py 我的定稿.docx 参考稿.pdf --baseline 我的原稿.docx

Prints:

  * prose overlap   -- mean share of 25-char shingles in my paragraphs that also
                       occur in the reference (paragraphs >= 60 normalised chars)
  * verbatim strings -- sentences of mine (>= 25 chars) that appear word for word
                       in the reference
  * shared headings -- short strings present in both

`--baseline` is not optional in practice. It runs the same measurement on the
draft you started from, so you can tell whether a number came from your edits or
was already there. Without it the report cannot be interpreted.
"""
import argparse
import io
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

try:
    from docx import Document
except ImportError:
    # Offline / locked-down Python: stdlib reader, same surface.
    from docxfallback import Document


def norm(t):
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]", "", t.lower())


def shingles(s, n):
    return {s[i:i + n] for i in range(max(0, len(s) - n + 1))}


def load(path):
    """Return (paragraphs, whole_text). PDFs give paragraphs=[]."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        if Document is None:
            sys.exit("python-docx is required for .docx")
        d = Document(path)
        paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
        cells = []
        for t in d.tables:
            for r in t.rows:
                for c in r.cells:
                    if c.text.strip():
                        cells.append(c.text.strip())
        return paras, "\n".join(paras + cells)
    if ext == ".pdf":
        from pdftext import extract_text
        return [], extract_text(path)
    if ext in (".txt", ".md"):
        return [], io.open(path, encoding="utf-8", errors="replace").read()
    sys.exit("unsupported file type: %s" % ext)


def prose(paras, min_chars):
    return [p for p in paras if len(norm(p)) >= min_chars]


def cover(text, ref_sh, n):
    s = shingles(norm(text), n)
    return (100.0 * len(s & ref_sh) / len(s)) if s else 0.0


def report(mine_paras, ref_blob, base_paras, n, min_chars, label):
    ref_norm = norm(ref_blob)
    ref_sh = shingles(ref_norm, n)
    base_norm = norm("\n".join(base_paras)) if base_paras else None
    print("reference corpus: %d normalised chars, %d shingles (n=%d)"
          % (len(ref_norm), len(ref_sh), n))

    base_mean = None
    if base_paras:
        ps = prose(base_paras, min_chars)
        covs = [cover(p, ref_sh, n) for p in ps]
        if covs:
            base_mean = sum(covs) / len(covs)
            print("BASELINE draft  : mean overlap %5.1f%% over %d prose paragraphs"
                  " (max %5.1f%%)" % (base_mean, len(ps), max(covs)))

    ps = prose(mine_paras, min_chars)
    rows = sorted(((cover(p, ref_sh, n), p) for p in ps), key=lambda r: -r[0])
    covs = [r[0] for r in rows]
    mean = sum(covs) / len(covs) if covs else 0.0
    print("%-16s: mean overlap %5.1f%% over %d prose paragraphs (max %5.1f%%)"
          % (label, mean, len(rows), covs[0] if covs else 0.0))

    print("\n-- paragraphs, worst first --")
    for c, p in rows:
        flag = "  <== HIGH" if c >= 40 else ""
        print("%5.1f%%  %s%s" % (c, p.replace("\n", " ")[:88], flag))

    def inherited(s):
        return base_norm is not None and norm(s) in base_norm

    print("\n-- verbatim strings (>=25 chars) also present in the reference --")
    hits = []
    for p in mine_paras:
        for sent in re.split(r"(?<=[.;:!?])\s+|(?<=[。；！？])", p):
            s = sent.strip()
            if len(norm(s)) >= 25 and norm(s) in ref_norm:
                hits.append(s)
    hits = list(dict.fromkeys(hits))
    for h in hits:
        tag = ("[inherited from draft]" if inherited(h)
               else "[NEW - came from the reference]  <== rewrite this")
        print("   -", h.replace("\n", " ")[:90], tag)
    if not hits:
        print("   (none)")

    print("\n-- headings / labels present in both --")
    heads = [p for p in mine_paras
             if len(p) < 70 and 8 <= len(norm(p)) < 60 and norm(p) in ref_norm]
    heads = list(dict.fromkeys(heads))
    for h in heads:
        tag = ("[inherited]" if inherited(h) else "[NEW]  <== rename")
        print("   +", h[:70], tag)
    if not heads:
        print("   (none)")

    new_hits = [h for h in hits if not inherited(h)]
    new_heads = [h for h in heads if not inherited(h)]

    print("\n-- verdict --")
    if base_mean is not None:
        print("    prose overlap %.1f%% vs baseline %.1f%%  (delta %+.1f pp) -> %s"
              % (mean, base_mean, mean - base_mean,
                 "your edits added no overlap" if abs(mean - base_mean) < 1.0
                 else ("independent" if mean <= 1 else
                       ("suspicious" if mean <= 15 else "TOO HIGH"))))
    else:
        print("    prose overlap %.1f%% -> %s   (no --baseline: cannot attribute)"
              % (mean, "independent" if mean <= 1 else
                 ("suspicious" if mean <= 15 else "TOO HIGH")))
    print("    verbatim strings: %d total, %d inherited, %d NEW"
          % (len(hits), len(hits) - len(new_hits), len(new_hits)))
    print("    shared headings:  %d total, %d NEW"
          % (len(heads), len(new_heads)))
    print("\n    -> %s" % ("PASS" if not new_hits and not new_heads
                           else "rewrite the NEW items above, then re-run"))


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser()
    ap.add_argument("mine", help="your document (.docx)")
    ap.add_argument("reference", help="the other document (.docx/.pdf/.txt)")
    ap.add_argument("--baseline", default=None,
                    help="the draft you started from; makes the delta meaningful")
    ap.add_argument("--n", type=int, default=25, help="shingle length in chars")
    ap.add_argument("--min-chars", type=int, default=60,
                    help="ignore paragraphs shorter than this when averaging")
    a = ap.parse_args()

    mine_paras, _blob = load(a.mine)
    if not mine_paras:
        sys.exit("the first argument must be a .docx - paragraph granularity is required")
    _r, ref_blob = load(a.reference)
    base_paras = load(a.baseline)[0] if a.baseline else None

    report(mine_paras, ref_blob, base_paras, a.n, a.min_chars,
           os.path.basename(a.mine))
