"""Minimal, dependency-free stand-in for the bits of python-docx these
scripts use.

Why this exists: python-docx is a third-party package. On a machine with no
network (or a locked-down Python), `pip install python-docx` fails and all
three scripts become unusable -- `flavor_audit.py` degrades to plain text
(layer 6 goes n/a) and `consistency.py` / `overlap.py` refuse to run at all,
which is exactly when a table-heavy document most needs checking.

The surface implemented here is deliberately small:

    Document(path)      -> .paragraphs, .tables, .inline_shapes, .element.body
    qn("w:drawing")     -> "{namespace}drawing"

Paragraph: .text, .runs            Run: .text, .bold
Table:     .rows                   Row: .cells        Cell: .text, .paragraphs

Text is taken from <w:t> only, so OMML equation text (<m:t>) is excluded --
the same thing python-docx does, which keeps prose statistics comparable.

Usage in the scripts:

    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError:
        from docxfallback import Document, qn
"""

import zipfile
import xml.etree.ElementTree as ET

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def qn(tag):
    """'w:drawing' -> '{ns}drawing'."""
    prefix, _, local = tag.partition(":")
    if prefix != "w":
        raise ValueError("docxfallback only knows the w: namespace: %r" % tag)
    return "{%s}%s" % (W, local)


def _w(local):
    return "{%s}%s" % (W, local)


def _text_of(el):
    """Concatenate every <w:t> below el, in document order."""
    return "".join(t.text or "" for t in el.iter(_w("t")))


class Run(object):
    def __init__(self, el):
        self._el = el

    @property
    def text(self):
        return _text_of(self._el)

    @property
    def bold(self):
        rpr = self._el.find(_w("rPr"))
        if rpr is None:
            return None
        b = rpr.find(_w("b"))
        if b is None:
            return None
        val = b.get(_w("val"))
        return val not in ("0", "false", "off")


class Paragraph(object):
    def __init__(self, el):
        self._el = el

    @property
    def text(self):
        return _text_of(self._el)

    @property
    def runs(self):
        return [Run(r) for r in self._el.findall(_w("r"))]

    @property
    def style(self):
        ppr = self._el.find(_w("pPr"))
        if ppr is None:
            return None
        st = ppr.find(_w("pStyle"))
        return st.get(_w("val")) if st is not None else None


class Cell(object):
    def __init__(self, el):
        self._el = el

    @property
    def text(self):
        return "\n".join(_text_of(p) for p in self._el.findall(_w("p")))

    @property
    def paragraphs(self):
        return [Paragraph(p) for p in self._el.findall(_w("p"))]


class Row(object):
    def __init__(self, el):
        self._el = el

    @property
    def cells(self):
        out = []
        for tc in self._el.findall(_w("tc")):
            cell = Cell(tc)
            span = 1
            tcpr = tc.find(_w("tcPr"))
            if tcpr is not None:
                gs = tcpr.find(_w("gridSpan"))
                if gs is not None:
                    try:
                        span = max(1, int(gs.get(_w("val"))))
                    except (TypeError, ValueError):
                        span = 1
            out.extend([cell] * span)   # python-docx repeats merged cells
        return out


class Table(object):
    def __init__(self, el):
        self._el = el

    @property
    def rows(self):
        return [Row(tr) for tr in self._el.findall(_w("tr"))]


class _ElemProxy(object):
    """Wraps an ElementTree element so the lxml-only calls the scripts make
    keep working: .getnext(), .iterchildren(), and .itertext() (delegated)."""

    def __init__(self, el, parent=None):
        self._el = el
        self._parent = parent

    def __getattr__(self, name):
        return getattr(self._el, name)

    def getnext(self):
        if self._parent is None:
            return None
        siblings = list(self._parent._el)
        try:
            i = siblings.index(self._el)
        except ValueError:
            return None
        if i + 1 >= len(siblings):
            return None
        return _ElemProxy(siblings[i + 1], self._parent)

    def iterchildren(self, tag=None):
        return [_ElemProxy(c, self) for c in list(self._el)
                if tag is None or c.tag == tag]

    def findall(self, *a, **kw):
        return [_ElemProxy(e, self) for e in self._el.findall(*a, **kw)]

    def iter(self, tag=None):
        return self._el.iter(tag)


class _Element(object):
    """Mimics python-docx's Document.element (only .body is used)."""

    def __init__(self, body):
        self.body = _ElemProxy(body)


class Document(object):
    def __init__(self, path):
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml")
        root = ET.fromstring(xml)
        body = root.find(_w("body"))
        if body is None:
            raise ValueError("no w:body in %s" % path)
        self._body = body
        self.element = _Element(body)

    @property
    def paragraphs(self):
        return [Paragraph(p) for p in self._body.findall(_w("p"))]

    @property
    def tables(self):
        return [Table(t) for t in self._body.findall(_w("tbl"))]

    @property
    def inline_shapes(self):
        return list(self._body.iter(_w("drawing")))


def fonts_in_use(doc):
    """Counter of explicit w:ascii families in the body (for the font-mixing
    check). Returns an empty Counter when the document carries none."""
    from collections import Counter
    counts = Counter()
    try:
        for el in doc.element.body.iter(qn("w:rFonts")):
            name = el.get(qn("w:ascii"))
            if name:
                counts[name] += 1
    except Exception:
        return Counter()
    return counts
